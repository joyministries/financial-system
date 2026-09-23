from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from sqlalchemy import and_, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.grade import Grade, Student
from app.models.invoice import Invoice
from app.models.payment import Payment
from app.models.schedule import AdditionalCharge
from app.services.ledger import LedgerService


class ReportService:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.ledger = LedgerService(db)

    async def monthly_income(self, academic_year: int, month: int) -> dict:
        start, end = self._month_range(academic_year, month)

        stmt = select(func.sum(Payment.amount)).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        result = await self.db.execute(stmt)
        total = result.scalar() or Decimal("0")

        stmt_count = select(func.count(Payment.id)).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        count_result = await self.db.execute(stmt_count)
        count = count_result.scalar() or 0

        return {
            "period": f"{academic_year}-{month:02d}",
            "total_income": str(total),
            "payment_count": count,
        }

    async def monthly_summary(
        self,
        academic_year: int,
        month: int,
        grade_id: str | None = None,
    ) -> dict:
        """Monthly admin dashboard view.

        Income actually received in the month combined with the Excel-ledger
        outstanding position up to (and including) that month — how much was
        collected, how much is still owed, and which students owe.
        """
        start, end = self._month_range(academic_year, month)

        # Income received during the month
        income_stmt = select(func.sum(Payment.amount)).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        count_stmt = select(func.count(Payment.id)).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        if grade_id:
            income_stmt = income_stmt.join(
                Student, Student.id == Payment.student_id
            ).where(Student.grade_id == grade_id)
            count_stmt = count_stmt.join(
                Student, Student.id == Payment.student_id
            ).where(Student.grade_id == grade_id)
        total_income = (await self.db.execute(income_stmt)).scalar() or Decimal("0")
        payment_count = (await self.db.execute(count_stmt)).scalar() or 0

        # Outstanding position from the Excel-aligned ledger
        rows = await self.ledger.students_outstanding(
            academic_year, grade_id=grade_id, up_to_month=month
        )
        owing = [r for r in rows if r["outstanding"] > 0]
        owing.sort(key=lambda r: r["outstanding"], reverse=True)

        return {
            "academic_year": academic_year,
            "month": month,
            "total_income": str(total_income),
            "payment_count": payment_count,
            "outstanding_total": str(
                sum((r["outstanding"] for r in owing), Decimal("0"))
            ),
            "students_owing": len(owing),
            "students_owing_list": [
                {
                    "student_id": r["student_id"],
                    "student_number": r["student_number"],
                    "name": r["name"],
                    "grade": r["grade"],
                    "balance": str(r["outstanding"]),
                }
                for r in owing
            ],
        }

    async def yearly_income(self, academic_year: int) -> dict:
        start = datetime(academic_year, 1, 1, tzinfo=UTC)
        end = datetime(academic_year + 1, 1, 1, tzinfo=UTC)

        stmt = select(Payment).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        result = await self.db.execute(stmt)
        payments = result.scalars().all()

        monthly_breakdown: dict[int, Decimal] = {}
        monthly_counts: dict[int, int] = {}
        for p in payments:
            m = p.payment_date.month
            monthly_breakdown[m] = monthly_breakdown.get(m, Decimal("0")) + p.amount
            monthly_counts[m] = monthly_counts.get(m, 0) + 1

        monthly_data = [
            {
                "period": f"{academic_year}-{m:02d}",
                "total_income": str(monthly_breakdown.get(m, Decimal("0"))),
                "payment_count": monthly_counts.get(m, 0),
            }
            for m in range(1, 13)
        ]

        total = sum(Decimal(d["total_income"]) for d in monthly_data)
        return {
            "academic_year": academic_year,
            "total_income": str(total),
            "monthly_breakdown": monthly_data,
        }

    async def outstanding_fees(self, academic_year: int) -> dict:
        rows = await self.ledger.students_outstanding(academic_year)
        owing = [r for r in rows if r["outstanding"] > 0]
        return {
            "academic_year": academic_year,
            "students_with_outstanding": len(owing),
            "students": [
                {
                    "student_id": r["student_id"],
                    "student_number": r["student_number"],
                    "name": r["name"],
                    "outstanding": str(r["outstanding"]),
                }
                for r in owing
            ],
        }

    async def payments_received(
        self,
        academic_year: int,
        grade_id: str | None = None,
        payment_method: str | None = None,
        month: int | None = None,
    ) -> dict:
        """Payments received, optionally scoped to a single month.

        When *month* is given every figure (totals, breakdown, by-method) is
        restricted to that month; otherwise the whole academic year is used.
        """
        if month is not None:
            start, end = self._month_range(academic_year, month)
        else:
            start = datetime(academic_year, 1, 1, tzinfo=UTC)
            end = datetime(academic_year + 1, 1, 1, tzinfo=UTC)

        stmt = select(Payment).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        if grade_id:
            stmt = stmt.join(
                Student, Student.id == Payment.student_id
            ).where(Student.grade_id == grade_id)
        if payment_method:
            stmt = stmt.where(Payment.payment_method == payment_method)
        result = await self.db.execute(stmt)
        payments = result.scalars().all()

        monthly_breakdown: dict[int, Decimal] = {}
        monthly_counts: dict[int, int] = {}
        by_method: dict[str, Decimal] = {}
        for p in payments:
            m = p.payment_date.month
            monthly_breakdown[m] = monthly_breakdown.get(m, Decimal("0")) + p.amount
            monthly_counts[m] = monthly_counts.get(m, 0) + 1
            method = p.payment_method or "unknown"
            by_method[method] = by_method.get(method, Decimal("0")) + p.amount

        return {
            "academic_year": academic_year,
            "month": month,
            "total_payments": str(sum(monthly_breakdown.values(), Decimal("0"))),
            "total_received": str(sum(monthly_breakdown.values(), Decimal("0"))),
            "payment_count": sum(monthly_counts.values()),
            "monthly_breakdown": [
                {
                    "period": f"{academic_year}-{m:02d}",
                    "total": str(monthly_breakdown.get(m, Decimal("0"))),
                    "count": monthly_counts.get(m, 0),
                }
                for m in (sorted(monthly_breakdown) if month is not None else range(1, 13))
            ],
            "by_method": {k: str(v) for k, v in sorted(by_method.items())},
        }

    async def payment_trends(self, academic_year: int) -> dict:
        start = datetime(academic_year, 1, 1, tzinfo=UTC)
        end = datetime(academic_year + 1, 1, 1, tzinfo=UTC)

        stmt = select(Payment).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        result = await self.db.execute(stmt)
        payments = result.scalars().all()

        monthly_data = {}
        for p in payments:
            m = p.payment_date.month
            entry = monthly_data.setdefault(m, {"total": Decimal("0"), "count": 0})
            entry["total"] += p.amount
            entry["count"] += 1

        return {
            "academic_year": academic_year,
            "months": [
                {
                    "month": m,
                    "total": str(monthly_data.get(m, {"total": Decimal("0")})["total"]),
                    "count": monthly_data.get(m, {"count": 0})["count"],
                }
                for m in range(1, 13)
            ],
        }

    async def statement_report(
        self,
        academic_year: int,
        status_filter: str | None = None,
        grade_id: str | None = None,
        month: int | None = None,
        month_only: bool = False,
    ) -> dict:
        """School-wide statement summary — every approved student with their
        outstanding balance from the Excel-aligned ledger, scoped to a selected
        month when supplied."""
        if month_only and month is not None:
            rows = await self._monthly_statement_rows(academic_year, month, grade_id)
        else:
            rows = await self.ledger.students_outstanding(
                academic_year, grade_id=grade_id, up_to_month=month
            )

        students = []
        total_outstanding = Decimal("0")
        for r in rows:
            balance = r["outstanding"]
            student_status = "paid" if balance <= 0 else "overdue"
            if student_status == "overdue":
                total_outstanding += balance

            if status_filter and student_status != status_filter:
                continue

            students.append({
                "student_id": r["student_id"],
                "student_number": r["student_number"],
                "name": r["name"],
                "grade": r["grade"] or "",
                "balance": str(balance),
                "status": student_status,
            })

        return {
            "academic_year": academic_year,
            "month": month,
            "month_only": month_only,
            "total_students": len(students),
            "total_outstanding": str(total_outstanding),
            "students": students,
        }

    async def _monthly_statement_rows(
        self, academic_year: int, month: int, grade_id: str | None = None
    ) -> list[dict]:
        """Approved students with only this month's billed/paid balance."""
        start, end = self._month_range(academic_year, month)

        inv_stmt = (
            select(
                Student.id,
                Student.student_number,
                Student.first_name,
                Student.last_name,
                Grade.name.label("grade_name"),
                func.coalesce(func.sum(Invoice.subtotal), 0).label("required"),
            )
            .join(Grade, Grade.id == Student.grade_id)
            .outerjoin(
                Invoice,
                and_(
                    Invoice.student_id == Student.id,
                    Invoice.status != "void",
                    Invoice.academic_year == academic_year,
                    Invoice.month == month,
                ),
            )
            .where(Student.registration_status == "approved")
        )
        if grade_id:
            inv_stmt = inv_stmt.where(Student.grade_id == grade_id)
        inv_stmt = inv_stmt.group_by(
            Student.id, Student.student_number, Student.first_name,
            Student.last_name, Grade.name,
        )
        inv_rows = (await self.db.execute(inv_stmt)).all()

        pay_stmt = (
            select(Student.id, func.coalesce(func.sum(Payment.amount), 0))
            .join(
                Payment,
                and_(
                    Payment.student_id == Student.id,
                    Payment.status == "verified",
                    Payment.payment_date >= start,
                    Payment.payment_date < end,
                ),
            )
            .where(Student.registration_status == "approved")
            .group_by(Student.id)
        )
        if grade_id:
            pay_stmt = pay_stmt.where(Student.grade_id == grade_id)
        paid_by_student = dict((await self.db.execute(pay_stmt)).all())

        charge_stmt = (
            select(Student.id, func.coalesce(func.sum(AdditionalCharge.amount), 0))
            .join(
                AdditionalCharge,
                and_(
                    AdditionalCharge.student_id == Student.id,
                    AdditionalCharge.academic_year == academic_year,
                    AdditionalCharge.month == month,
                ),
            )
            .where(Student.registration_status == "approved")
            .group_by(Student.id)
        )
        if grade_id:
            charge_stmt = charge_stmt.where(Student.grade_id == grade_id)
        charges_by_student = dict((await self.db.execute(charge_stmt)).all())

        return [
            {
                "student_id": r.id,
                "student_number": r.student_number,
                "name": f"{r.first_name} {r.last_name}",
                "grade": r.grade_name,
                "required": Decimal(str(r.required)) + charges_by_student.get(r.id, Decimal("0")),
                "paid": paid_by_student.get(r.id, Decimal("0")),
                "outstanding": (
                    Decimal(str(r.required))
                    + charges_by_student.get(r.id, Decimal("0"))
                    - paid_by_student.get(r.id, Decimal("0"))
                ),
            }
            for r in inv_rows
        ]

    async def students_xlsx(
        self,
        academic_year: int,
        grade_id: str | None = None,
        month: int | None = None,
    ) -> BytesIO:
        """Admin Excel export in the school's suspension-list layout.

        Produces an .xlsx workbook with the same shape as the office's
        "LCS GERMISTON <MONTH> SUSPENSION LIST" spreadsheet (yellow bold
        headers: Customer | Grade | Amount | Comments | Learners on
        suspension). Every approved student is included once, optionally
        scoped to a single grade.

        When `month` is given the Amount column is the outstanding balance
        FOR that month (invoices with month <= month minus verified payments
        received up to the end of that month) from the Excel-aligned ledger —
        the same position the month's statement shows. When `month` is omitted
        the Amount is the full-year outstanding. Either way a SUM footer row
        is appended.

        Grade cells mirror the office convention: 'RR' / 'R' or 1-9.
        """
        rows = await self.ledger.students_outstanding(
            academic_year, grade_id=grade_id, up_to_month=month
        )

        def _grade_display(name: str) -> str:
            if name == "GRADE R":
                return "R"
            if name == "GRADE RR":
                return "RR"
            return name.replace("GRADE ", "", 1)

        sheet_label = (
            datetime(academic_year, month, 1).strftime("%B").upper()
            if month else "STUDENTS"
        )
        headers = ["Customer", "Grade", "Amount", "Comments", "Learners on suspension"]

        surfacing_rows = [
            [
                f"({r['student_number']}) {r['name']}".strip(),
                _grade_display(r["grade"]),
                float(r["outstanding"]),
                None,
                None,
            ]
            for r in rows
        ]

        # Imported lazily so the REST API continues to boot even if the
        # spreadsheet library is missing/mis-installed in the serverless runtime.
        # If it is unavailable, fall back to the stdlib writer so the export
        # endpoint keeps working regardless of what the runtime bundled.
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            from app.services.xlsx_writer import build_suspension_list_xlsx

            sum_total = sum(row[2] for row in surfacing_rows)
            return build_suspension_list_xlsx(
                sheet_label=sheet_label,
                headers=headers,
                data_rows=surfacing_rows,
                column_widths=(
                    ("A", 56.89), ("B", 28.66), ("C", 21.66),
                    ("D", 14.33), ("E", 24.11),
                ),
                sum_footer=bool(surfacing_rows),
                sum_total=sum_total,
            )

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_label

        header_fill = PatternFill("solid", fgColor="FFFF00")
        header_font = Font(bold=True)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        ws["A1"].alignment = Alignment(horizontal="left")
        ws["B1"].alignment = Alignment(horizontal="center")
        ws["C1"].alignment = Alignment(horizontal="center")
        ws["E1"].alignment = Alignment(horizontal="center")

        for row_values in surfacing_rows:
            ws.append(row_values)

        if len(rows) > 0:
            ws.append([None, "OUTSTANDING BALANCE", f"=SUM(C2:C{len(rows) + 1})", None, None])

        for col_letter, width in (
            ("A", 56.89), ("B", 28.66), ("C", 21.66),
            ("D", 14.33), ("E", 24.11),
        ):
            ws.column_dimensions[col_letter].width = width

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    async def carry_forward(self, academic_year: int, month: int) -> dict:
        """Dashboard carry-forward section.

        - not_paid: active students with NO verified payment in the selected
          month (they have not paid for that month yet)
        - outstanding: students with a positive ledger balance in the academic
          year up to (and including) the selected month
        """
        start, end = self._month_range(academic_year, month)

        paid_stmt = select(Payment.student_id).where(
            Payment.status == "verified",
            Payment.payment_date >= start,
            Payment.payment_date < end,
        )
        paid_result = await self.db.execute(paid_stmt)
        paid_ids = set(paid_result.scalars().all())

        stmt = (
            select(Student, Grade.name)
            .join(Grade, Grade.id == Student.grade_id)
            .where(Student.is_active == True)  # noqa: E712
            .order_by(Student.last_name)
        )
        result = await self.db.execute(stmt)
        rows = result.all()

        not_paid = [
            {
                "student_id": s.id,
                "student_number": s.student_number,
                "name": f"{s.first_name} {s.last_name}",
                "grade": grade_name,
            }
            for s, grade_name in rows
            if s.id not in paid_ids
        ]

        ledger_rows = await self.ledger.students_outstanding(
            academic_year, up_to_month=month
        )
        outstanding = [
            {
                "student_id": r["student_id"],
                "student_number": r["student_number"],
                "name": r["name"],
                "grade": r["grade"],
                "balance": str(r["outstanding"]),
            }
            for r in ledger_rows
            if r["outstanding"] > 0
        ]
        outstanding.sort(key=lambda r: Decimal(r["balance"]), reverse=True)

        return {
            "academic_year": academic_year,
            "month": month,
            "not_paid_count": len(not_paid),
            "not_paid": not_paid,
            "outstanding_count": len(outstanding),
            "outstanding": outstanding,
        }

    def _month_range(self, year: int, month: int) -> tuple[datetime, datetime]:
        start = datetime(year, month, 1, tzinfo=UTC)
        if month == 12:
            end = datetime(year + 1, 1, 1, tzinfo=UTC)
        else:
            end = datetime(year, month + 1, 1, tzinfo=UTC)
        return start, end
