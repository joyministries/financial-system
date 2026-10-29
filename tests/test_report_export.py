from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest import mock

import pytest
from openpyxl import load_workbook

from app.services import report as report_mod
from app.services.report import ReportService
from app.services.xlsx_writer import build_suspension_list_xlsx

HEADERS = [
    "Customer", "Grade", "Amount", "Fees for remaining months (R)",
    "Comments", "Learners on suspension",
]

FOOTER_LABEL = "TOTAL AMOUNT DUE"


def test_fallback_writer_builds_valid_workbook():
    buf = build_suspension_list_xlsx(
        sheet_label="SEPTEMBER",
        headers=HEADERS,
        data_rows=[
            ["(1001) Alex Kuti", "8", 1250.5, 3751.5, None, None],
            ["(1002) Thandi Kunene", "R", 0.0, 0.0, None, None],
            ["(1003) O'Brien & Son", "1", 499.99, 1499.97, None, None],
        ],
        column_widths=(("A", 56.89), ("B", 28.66), ("C", 21.66),
                       ("D", 21.66), ("E", 14.33), ("F", 24.11)),
        footer=[None, FOOTER_LABEL, 1750.49, 5251.47, None, None],
    )
    buf.seek(0)
    wb = load_workbook(BytesIO(buf.read()))

    ws = wb.active
    assert ws.title == "SEPTEMBER"

    # Header row styling
    assert [c.value for c in ws[1]] == HEADERS
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb == "FFFFFF00"
    # A left, B/C/E center (matches the openpyxl version)
    assert ws["A1"].alignment.horizontal == "left"
    assert ws["B1"].alignment.horizontal == "center"
    assert ws["C1"].alignment.horizontal == "center"
    assert ws["E1"].alignment.horizontal == "center"

    # Data rows
    assert ws["A2"].value == "(1001) Alex Kuti"
    assert ws["B2"].value == "8"
    assert ws["C2"].value == 1250.5
    assert ws["A3"].value == "(1002) Thandi Kunene"
    assert ws["C3"].value == 0.0
    # Ampersand + apostrophe survive the XML escaping round-trip
    assert ws["A4"].value == "(1003) O'Brien & Son"

    # Totals footer: static figures, not =SUM(), so engine-less previewers
    # (Drive, phone viewers) still show the numbers.
    assert ws["B5"].value == FOOTER_LABEL
    assert ws["C5"].value == 1750.49
    assert ws["D5"].value == 5251.47

    # Column widths
    assert ws.column_dimensions["A"].width == 56.89


def _student(student_id, number, first, last, grade_name, is_active=True):
    """A (Student, grade_name) pair as students_monthly_fee selects it."""
    return (
        SimpleNamespace(
            id=student_id,
            student_number=number,
            first_name=first,
            last_name=last,
            is_active=is_active,
        ),
        grade_name,
    )


def _fake_fee_db(pairs):
    """Emulate the single query ReportService.students_monthly_fee runs."""
    db = SimpleNamespace()
    db.execute = mock.AsyncMock(return_value=SimpleNamespace(all=lambda: list(pairs)))
    return db


def _fake_fees(**by_student):
    """Patch load_monthly_fee_lookup to hand back fixed monthly instalments."""
    async def lookup(_db, students, _year):
        return {s.id: by_student.get(s.id, Decimal("0")) for s in students}

    return mock.patch.object(report_mod, "load_monthly_fee_lookup", lookup)


def _sheet(buf):
    buf.seek(0)
    return load_workbook(BytesIO(buf.read())).active


def _cells(ws, cols=6):
    return [
        [c.value for c in row]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=cols)
    ]


@pytest.mark.asyncio
async def test_students_xlsx_falls_back_when_openpyxl_missing():
    """Exercise the real ReportService.students_xlsx() with openpyxl
    blocked at import time — the exact deployed failure mode — and
    assert the stdlib fallback produces a loadable workbook.
    """
    fake_db = _fake_fee_db([
        _student("s1", "1001", "Alex", "Kuti", "GRADE 8"),
        _student("s2", "1002", "Thandi", "Kunene", "GRADE R"),
    ])

    real_import = __import__

    def _block_openpyxl(name, globals=None, locals=None, fromlist=(), level=0):
        if name == "openpyxl":
            raise ImportError("No module named 'openpyxl'")
        return real_import(name, globals, locals, fromlist, level)

    service = ReportService(db=fake_db)  # type: ignore[arg-type]

    with _fake_fees(s1=Decimal("2980.00"), s2=Decimal("1700.00")), \
            mock.patch("builtins.__import__", side_effect=_block_openpyxl):
        buf = await service.students_xlsx(academic_year=2026, month=9)

    ws = _sheet(buf)
    assert ws.title == "SEPTEMBER"
    assert [c.value for c in ws[1]] == HEADERS
    assert ws["A2"].value == "(1001) Alex Kuti"
    assert ws["B2"].value == "8"  # GRADE 8 -> "8"
    assert ws["B3"].value == "R"  # GRADE R -> "R"
    # September: the month's instalment, then Oct+Nov+Dec still to come.
    assert ws["C2"].value == 2980.0
    assert ws["D2"].value == 8940.0
    assert ws["C3"].value == 1700.0
    assert ws["D3"].value == 5100.0
    assert ws["B4"].value == FOOTER_LABEL
    assert ws["C4"].value == 4680.0
    assert ws["D4"].value == 14040.0


@pytest.mark.asyncio
async def test_students_xlsx_matches_the_fallback_cell_for_cell():
    """Parity check: the primary openpyxl path and the stdlib fallback must
    emit the identical sheet, whichever one the runtime can load.
    """
    pairs = [
        _student("s1", "1001", "Alex", "Kuti", "GRADE 8"),
        _student("s2", "1002", "Thandi", "Kunene", "GRADE RR"),
    ]
    fees = dict(s1=Decimal("2980.00"), s2=Decimal("1700.00"))

    real_import = __import__

    def _block_openpyxl(name, globals=None, locals=None, fromlist=(), level=0):
        if name == "openpyxl":
            raise ImportError("No module named 'openpyxl'")
        return real_import(name, globals, locals, fromlist, level)

    with _fake_fees(**fees):
        primary = _sheet(
            await ReportService(db=_fake_fee_db(pairs)).students_xlsx(2026, month=9)
        )
        with mock.patch("builtins.__import__", side_effect=_block_openpyxl):
            fallback = _sheet(
                await ReportService(db=_fake_fee_db(pairs)).students_xlsx(2026, month=9)
            )

    assert primary.title == fallback.title == "SEPTEMBER"
    assert _cells(primary) == _cells(fallback)
    assert _cells(primary)[0] == HEADERS
    assert _cells(primary)[-1] == [None, FOOTER_LABEL, 4680.0, 14040.0, None, None]


# ---------------------------------------------------------------------------
# Per-student x per-month outstanding matrix (the Reports -> Outstanding Fees
# Excel/CSV export). The export needs one row per student with a column per
# month instead of a single cumulative figure for the whole year.
# ---------------------------------------------------------------------------


def _matrix_row(student_id, student_number, first_name, last_name, grade, outstanding):
    """A row shaped like LedgerService.students_outstanding."""
    return {
        "student_id": student_id,
        "student_number": student_number,
        "name": f"{first_name} {last_name}",
        "grade": grade,
        "required": Decimal(outstanding),
        "paid": Decimal("0"),
        "outstanding": Decimal(outstanding),
    }


class _FakeLedger:
    """Stand-in for LedgerService.students_outstanding.

    Returns the row set registered for whichever ``up_to_month`` was asked
    for, so the matrix's per-month columns can be asserted with no database.
    """

    def __init__(self, by_month):
        self.by_month = by_month
        self.calls = []

    async def students_outstanding(self, academic_year, grade_id=None, up_to_month=None):
        self.calls.append(up_to_month)
        return self.by_month.get(up_to_month, [])


@pytest.mark.asyncio
async def test_outstanding_matrix_is_running_balance_per_month():
    """Carry-over mode: every month is the balance at that month's end."""
    ledger = _FakeLedger({
        1: [
            _matrix_row("s1", "1001", "Alex", "Kuti", "GRADE 8", "1000"),
            _matrix_row("s2", "1002", "Thandi", "Kunene", "GRADE R", "500"),
        ],
        2: [
            _matrix_row("s1", "1001", "Alex", "Kuti", "GRADE 8", "400"),
            _matrix_row("s2", "1002", "Thandi", "Kunene", "GRADE R", "500"),
        ],
        # Thandi has no rows from March on; she must still get a full row.
        3: [_matrix_row("s1", "1001", "Alex", "Kuti", "GRADE 8", "400")],
    })
    service = ReportService(db=SimpleNamespace())
    service.ledger = ledger

    result = await service.outstanding_matrix(2026, month_only=False, up_to_month=3)

    assert ledger.calls == [1, 2, 3]
    assert result["month_only"] is False
    assert [m["label"] for m in result["months"]] == ["January", "February", "March"]
    assert [s["name"] for s in result["students"]] == ["Alex Kuti", "Thandi Kunene"]

    alex = result["students"][0]
    assert alex["student_number"] == "1001"
    assert alex["grade"] == "GRADE 8"
    assert alex["balances"] == {"1": "1000", "2": "400", "3": "400"}

    # Absent from March -> zero, never a missing key.
    assert result["students"][1]["balances"] == {"1": "500", "2": "500", "3": "0"}

    # Per-month column totals, which replace the old cumulative block.
    assert result["totals"] == {"1": "1500", "2": "900", "3": "400"}


@pytest.mark.asyncio
async def test_outstanding_matrix_month_only_uses_that_months_own_fee():
    """Fee mode: each column is that month's instalment, not a running total."""
    service = ReportService(db=SimpleNamespace())
    seen = []

    async def fake_fee(academic_year, month, grade_id=None):
        seen.append(month)
        return [_matrix_row("s1", "1001", "Alex", "Kuti", "GRADE 8", str(month * 100))]

    service.students_monthly_fee = fake_fee

    result = await service.outstanding_matrix(2026, month_only=True, up_to_month=4)

    assert seen == [1, 2, 3, 4]
    assert result["month_only"] is True
    assert result["students"][0]["balances"] == {
        "1": "100", "2": "200", "3": "300", "4": "400",
    }
    assert result["totals"] == {"1": "100", "2": "200", "3": "300", "4": "400"}


@pytest.mark.asyncio
async def test_outstanding_matrix_defaults_to_full_year():
    """Omitting up_to_month exports Jan..Dec, and a student seen only in a
    late month still appears with zeros for the earlier months."""
    ledger = _FakeLedger({
        11: [_matrix_row("s9", "1009", "Late", "Enrolment", "GRADE 1", "750")],
    })
    service = ReportService(db=SimpleNamespace())
    service.ledger = ledger

    result = await service.outstanding_matrix(2026)

    assert ledger.calls == list(range(1, 13))
    assert len(result["months"]) == 12
    assert result["months"][0]["label"] == "January"
    assert result["months"][11]["label"] == "December"

    balances = result["students"][0]["balances"]
    assert balances["1"] == "0"
    assert balances["11"] == "750"
    assert result["totals"]["11"] == "750"


# ---------------------------------------------------------------------------
# "Fee for this month" — the reports' month view (screen, /reports/statements,
# /reports/export-excel, /reports/export-students).
#
# The figure is the grade's monthly tuition instalment and nothing else: not
# the generated statement document, not the balance carried over from earlier
# months, and no credit for payments already received. The companion column
# prices that same instalment once for every month STILL TO COME after the
# report month, so October counts November and December, and December counts
# nothing.
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_fee_view_is_the_grade_instalment_only():
    service = ReportService(db=_fake_fee_db([
        _student("s1", "1001", "Alex", "Kuti", "GRADE 1"),
    ]))

    with _fake_fees(s1=Decimal("1940.00")):
        rows = await service.students_monthly_fee(2026, 10)

    assert len(rows) == 1
    assert rows[0]["outstanding"] == Decimal("1940.00")
    assert rows[0]["required"] == Decimal("1940.00")
    # November and December remain.
    assert rows[0]["future_fees"] == Decimal("3880.00")
    # Row shape shared with the ledger, minus the paid column this view ignores.
    assert set(rows[0]) == {
        "student_id", "student_number", "name", "grade",
        "required", "outstanding", "future_fees",
    }


@pytest.mark.asyncio
async def test_future_fees_counts_only_the_months_after_the_report_month():
    """October prices two months, November one, December none."""
    pairs = [_student("s1", "1001", "Alex", "Kuti", "GRADE 1")]
    expected = {10: "3880.00", 11: "1940.00", 12: "0", 1: "21340.00"}

    for month, future in expected.items():
        service = ReportService(db=_fake_fee_db(pairs))
        with _fake_fees(s1=Decimal("1940.00")):
            rows = await service.students_monthly_fee(2026, month)

        assert rows[0]["future_fees"] == Decimal(future), f"month {month}"
        assert rows[0]["outstanding"] == Decimal("1940.00"), "the month's own fee"


@pytest.mark.asyncio
async def test_fee_view_skips_deactivated_and_test_seeded_students():
    service = ReportService(db=_fake_fee_db([
        _student("s1", "1001", "Alex", "Kuti", "GRADE 1"),
        _student("s2", "1002", "Gone", "School", "GRADE 1", is_active=False),
        _student("s3", "1003", "Test", "Child", "GRADE 1"),
    ]))

    with _fake_fees(s1=Decimal("1940.00"), s2=Decimal("1940.00"), s3=Decimal("1940.00")):
        rows = await service.students_monthly_fee(2026, 10)

    assert [r["student_id"] for r in rows] == ["s1"]


@pytest.mark.asyncio
async def test_fee_view_keeps_students_with_a_null_is_active():
    """Imported learners can carry NULL is_active; they must not vanish."""
    service = ReportService(db=_fake_fee_db([
        _student("s1", "1001", "Alex", "Kuti", "GRADE 1", is_active=None),
    ]))

    with _fake_fees(s1=Decimal("1940.00")):
        rows = await service.students_monthly_fee(2026, 10)

    assert [r["student_id"] for r in rows] == ["s1"]


@pytest.mark.asyncio
async def test_statement_report_month_mode_lists_every_student_with_a_fee():
    """No payment credit means a paid-up learner still lists at their fee."""
    service = ReportService(db=_fake_fee_db([
        _student("s1", "1001", "Alex", "Kuti", "GRADE 1"),
        _student("s2", "1002", "Thandi", "Kunene", "GRADE R"),
    ]))

    with _fake_fees(s1=Decimal("1940.00"), s2=Decimal("1700.00")):
        result = await service.statement_report(2026, "overdue", None, 10, month_only=True)

    assert result["month_only"] is True
    assert result["total_students"] == 2
    assert result["total_outstanding"] == "3640.00"
    assert result["future_fees_total"] == "7280.00"
    assert result["students"][0]["balance"] == "1940.00"
    assert result["students"][0]["future_fees"] == "3880.00"


@pytest.mark.asyncio
async def test_statement_report_carry_mode_still_uses_the_ledger():
    """The opt-in carry-over view is untouched: it remains the running balance."""
    service = ReportService(db=SimpleNamespace())
    service.ledger = _FakeLedger({
        10: [_matrix_row("s1", "1001", "Alex", "Kuti", "GRADE 1", "8400")],
    })

    result = await service.statement_report(2026, "overdue", None, 10, month_only=False)

    assert service.ledger.calls == [10]
    assert result["students"][0]["balance"] == "8400"
    assert "future_fees" not in result["students"][0]
